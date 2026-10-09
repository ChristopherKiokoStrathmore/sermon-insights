import json
from pathlib import Path

from openpyxl import load_workbook

from sermon_insights.config import Config
from sermon_insights.pipeline import run_pipeline

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "samples" / "offline" / "fixture"


def test_offline_run_writes_five_sheets_without_network(tmp_path: Path, monkeypatch):
    def boom(*_args, **_kwargs):
        raise AssertionError("network call")

    monkeypatch.setattr("subprocess.run", boom)
    cfg = Config(
        channel="@example-fellowship",
        since=__import__("datetime").date(2025, 1, 1),
        languages=["sw", "en"],
        cache_dir=tmp_path / "cache",
        output=tmp_path / "out.xlsx",
        network=False,
        offline_fixture=FIXTURE,
        topics_backend="keyword",
        summary_backend="lexical",
        throttle_between_videos=0,
        throttle_requests=0,
        throttle_subtitles=0,
        throttle_ytsearch=0,
    )
    report = run_pipeline(cfg)
    assert report["catalog"]["videos"] == 2
    assert report["captions"]["skipped"] == 2 or report["captions"]["cached"] == 2
    book = load_workbook(tmp_path / "out.xlsx")
    assert book.sheetnames == [
        "Sermons",
        "Scripture",
        "Worship Songs",
        "Song List",
        "Topics & Verses",
    ]
    sermons = list(book["Sermons"].iter_rows(min_row=2, values_only=True))
    assert len(sermons) == 1
    assert sermons[0][0] == "2025-06-15"
    assert sermons[0][7] == "high"
    assert "1 Corinthians 13:4" in sermons[0][16]
    assert sermons[0][12] == "Love & Relationships"
    assert "keyword-seed" in sermons[0][19]
    refs = [row[2] for row in book["Scripture"].iter_rows(min_row=2, values_only=True)]
    assert "1 Corinthians 13:4" in refs
    assert any(str(ref).startswith("Psalms 23") for ref in refs)
    songs = list(book["Worship Songs"].iter_rows(min_row=2, values_only=True))
    titles = [row[6] for row in songs]
    assert "Uongezeke Yesu" in titles
    assert any(str(title).startswith("Nje ya Lango") and "(verify)" in str(title) for title in titles)
    listed = [row[0] for row in book["Song List"].iter_rows(min_row=2, values_only=True) if row[0]]
    assert "Uongezeke Yesu" in listed
    assert not any(title and "Nje ya Lango" in str(title) for title in listed)
    assert book["Topics & Verses"]["B2"].value == "Love & Relationships"
    again = run_pipeline(cfg)
    assert again["catalog"]["videos"] == 2
    raw = json.loads((tmp_path / "cache" / "raw" / "streams.json").read_text(encoding="utf-8"))
    assert any(entry["id"] == "FixSermon1a" for entry in raw["entries"])


def test_private_and_old_streams_stay_out_of_the_catalog(tmp_path: Path):
    cfg = Config(
        channel="@example-fellowship",
        since=__import__("datetime").date(2025, 1, 1),
        cache_dir=tmp_path / "cache",
        output=tmp_path / "out.xlsx",
        network=False,
        offline_fixture=FIXTURE,
        topics_backend="keyword",
        summary_backend="lexical",
    )
    run_pipeline(cfg, ["catalog"])
    import duckdb

    con = duckdb.connect(str(tmp_path / "cache" / "sermon.duckdb"))
    ids = {row[0] for row in con.execute("select id from catalog").fetchall()}
    con.close()
    assert ids == {"FixSermon1a", "FixWorship1"}
    assert "PrivateVid1" not in ids
    assert "OldSermon01" not in ids
