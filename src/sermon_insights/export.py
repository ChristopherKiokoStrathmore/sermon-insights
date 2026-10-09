"""Five-sheet workbook: Sermons, Scripture, Worship Songs, Song List, Topics & Verses."""

from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sermon_insights.scripture import display_ref, get_index
from sermon_insights.songs import song_rows

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THEME_FILL = PatternFill("solid", fgColor="D9E1F2")
LOW_FILL = PatternFill("solid", fgColor="FCE4D6")
CONF_FILL = {
    "high": PatternFill("solid", fgColor="E2EFDA"),
    "medium": PatternFill("solid", fgColor="FFF2CC"),
    "low": LOW_FILL,
}
LINK_FONT = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")

SERMON_HEADERS = [
    "Date",
    "Title",
    "URL",
    "Sermon start",
    "Sermon end",
    "Start link",
    "Detection method",
    "Detection confidence",
    "Transcript source",
    "Preacher",
    "Summary (extractive)",
    "Opening (original language)",
    "Broad theme",
    "Message",
    "Main passage",
    "Main passage source",
    "Scripture refs (high/medium)",
    "Low-confidence refs flagged",
    "Summary source",
    "Theme method",
]
SCRIPTURE_HEADERS = [
    "Stream date",
    "Title",
    "Reference",
    "Book",
    "Confidence",
    "Flag / reason",
    "Timestamp",
    "Link",
    "Matched text",
    "Transcript snippet",
]
WORSHIP_HEADERS = [
    "Stream date",
    "Stream title",
    "Song start",
    "Song end",
    "Link at start",
    "Lyric snippet heard",
    "Official song title",
    "Original artist",
    "Official YouTube link",
    "Match method",
    "Confidence",
    "Notes",
]
SONG_HEADERS = [
    "Official title",
    "Original artist",
    "Official YouTube link",
    "Times sung (streams)",
    "Dates sung",
    "Song segments",
]
TOPIC_HEADERS = [
    "Level",
    "Theme / sermon message",
    "Summary",
    "Sermon date(s)",
    "Bible verses quoted (canonical order; times quoted; link to first quote)",
    "Unique verses",
]


def hms(seconds) -> str:
    seconds = int(seconds or 0)
    return f"{seconds // 3600}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def as_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value or "")


def _style_sheet(worksheet, widths: list[int]) -> None:
    for cell in worksheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def _link(cell, url: str | None) -> None:
    if url and str(url).startswith("http"):
        cell.hyperlink = str(url)
        cell.font = LINK_FONT


def _sermon_rows(store) -> list[dict]:
    analysis = store.analysis()
    rows = []
    for video in store.catalog("sermon"):
        info = analysis.get(video["id"])
        span = store.segment(video["id"]) or {}
        if not info and span.get("sermon_start_s") is None:
            continue
        rows.append(
            {
                "video": video,
                "span": span,
                "refs": store.scripture_for(video["id"]),
                "info": info or {},
            }
        )
    rows.sort(key=lambda row: as_date(row["video"]["upload_date"]), reverse=True)
    return rows


def _verse_text(records: list[dict], index) -> tuple[str, int]:
    bucket: dict[str, dict] = {}
    for record in sorted(records, key=lambda row: as_date(row["video"]["upload_date"])):
        video = record["video"]
        info = record["info"]
        span = record["span"]
        start = span.get("sermon_start_s") or 0
        main = info.get("main_passage") or ""
        if main:
            for piece in [part.strip() for part in str(main).split(";") if part.strip()]:
                slot = bucket.setdefault(
                    piece,
                    {
                        "n": 0,
                        "link": f"{video['url']}&t={int(start)}s",
                        "low": False,
                        "main": set(),
                    },
                )
                slot["main"].add(info.get("main_source") or "main text")
        for ref in sorted(record["refs"], key=lambda item: item["ts"]):
            slot = bucket.setdefault(
                ref["ref"],
                {"n": 0, "link": ref["link"], "low": False, "main": set()},
            )
            slot["n"] += 1
            slot["low"] = slot["low"] or ref["conf"] == "low"
    lines = []
    for ref, slot in sorted(bucket.items(), key=lambda item: index.sort_key(item[0])):
        tag = ""
        if slot["main"]:
            tag += " [" + " / ".join(sorted(slot["main"])) + "]"
        if slot["low"]:
            tag += " (verify)"
        lines.append(f"{ref}{tag} - quoted {slot['n']}x - {slot['link']}")
    return "\n".join(lines), len(bucket)


def write_workbook(store, output: Path, *, languages: list[str] | None = None, export_csv: bool = True) -> Path:
    index = get_index(languages)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    sermons = _sermon_rows(store)
    songs = store.songs()
    cache = song_rows(store.songcache())

    workbook = Workbook()
    sermon_sheet = workbook.active
    sermon_sheet.title = "Sermons"
    sermon_sheet.append(SERMON_HEADERS)
    for record in sermons:
        video = record["video"]
        span = record["span"]
        info = record["info"]
        refs = record["refs"]
        strong = sorted(
            {ref["ref"] for ref in refs if ref["conf"] != "low"},
            key=index.sort_key,
        )
        start = span.get("sermon_start_s") or 0
        end = span.get("sermon_end_s") or video.get("duration_s") or 0
        sermon_sheet.append(
            [
                as_date(video["upload_date"]),
                video["title"],
                video["url"],
                hms(start),
                hms(end),
                f"{video['url']}&t={int(start)}s",
                span.get("method") or "",
                span.get("confidence") or "",
                span.get("source") or "",
                info.get("preacher") or "Not detected",
                info.get("summary_en") or "",
                info.get("summary_sw") or "",
                info.get("theme") or "",
                info.get("message") or "",
                info.get("main_passage") or "",
                info.get("main_source") or "",
                "; ".join(strong),
                sum(1 for ref in refs if ref["conf"] == "low"),
                info.get("summary_source") or "",
                info.get("theme_method") or "",
            ]
        )
    _style_sheet(sermon_sheet, [12, 36, 42, 12, 12, 42, 36, 14, 16, 22, 60, 50, 28, 40, 24, 28, 42, 12, 22, 36])
    for row in sermon_sheet.iter_rows(min_row=2):
        _link(row[2], row[2].value)
        _link(row[5], row[5].value)

    scripture_sheet = workbook.create_sheet("Scripture")
    scripture_sheet.append(SCRIPTURE_HEADERS)
    for record in sermons:
        video = record["video"]
        for ref in record["refs"]:
            scripture_sheet.append(
                [
                    as_date(video["upload_date"]),
                    video["title"],
                    display_ref(ref["ref"], ref["conf"]),
                    ref["book"],
                    ref["conf"],
                    ref.get("why") or "",
                    hms(ref["ts"]),
                    ref["link"],
                    ref.get("span") or "",
                    ref.get("snippet") or "",
                ]
            )
    _style_sheet(scripture_sheet, [12, 36, 28, 18, 12, 42, 12, 42, 36, 60])
    for row in scripture_sheet.iter_rows(min_row=2):
        _link(row[7], row[7].value)
        if row[4].value == "low":
            for cell in row:
                cell.fill = LOW_FILL

    worship_sheet = workbook.create_sheet("Worship Songs")
    worship_sheet.append(WORSHIP_HEADERS)
    for row in songs:
        worship_sheet.append(
            [
                as_date(row["upload_date"]),
                row["stream_title"],
                hms(row["start_s"]),
                hms(row["end_s"]),
                f"{row['url']}&t={int(row['start_s'])}s",
                row.get("snippet") or "",
                row.get("title") or "Unidentified",
                row.get("artist") or "",
                row.get("youtube") or "",
                row.get("method") or "",
                row.get("conf") or "",
                row.get("notes") or "",
            ]
        )
    _style_sheet(worship_sheet, [12, 36, 12, 12, 42, 50, 36, 28, 42, 55, 12, 50])
    for row in worship_sheet.iter_rows(min_row=2):
        _link(row[4], row[4].value)
        _link(row[8], row[8].value)
        fill = CONF_FILL.get(row[10].value or "")
        if fill:
            row[10].fill = fill

    song_sheet = workbook.create_sheet("Song List")
    song_sheet.append(SONG_HEADERS)
    grouped: dict[str, dict] = {}
    for row in songs:
        title = row.get("title") or "Unidentified"
        if title == "Unidentified" or row.get("conf") not in ("high", "medium"):
            continue
        key = re.sub(r"[^a-z]", "", re.sub(r"\(.*?\)", "", title.lower()))
        slot = grouped.setdefault(
            key,
            {
                "title": title,
                "artist": row.get("artist") or "",
                "link": row.get("youtube") or cache.get(title, {}).get("link", ""),
                "dates": [],
                "videos": set(),
                "segments": 0,
            },
        )
        slot["dates"].append(as_date(row["upload_date"]))
        slot["videos"].add(row["video_id"])
        slot["segments"] += 1
        if not slot["link"] and row.get("youtube"):
            slot["link"] = row["youtube"]
    ordered = sorted(grouped.values(), key=lambda item: (-len(item["videos"]), item["title"].lower()))
    for slot in ordered:
        song_sheet.append(
            [
                slot["title"],
                slot["artist"],
                slot["link"] or "(no single official upload)",
                len(slot["videos"]),
                ", ".join(sorted(set(slot["dates"]), reverse=True)),
                slot["segments"],
            ]
        )
    low_titled = sum(1 for row in songs if row.get("title") != "Unidentified" and row.get("conf") == "low")
    unidentified = sum(1 for row in songs if row.get("title") == "Unidentified")
    song_sheet.append([])
    song_sheet.append(
        [
            (
                f"Excluded: {low_titled + unidentified} segments "
                f"({low_titled} titled but low-confidence / verify, {unidentified} unidentified) "
                f"out of {len(songs)} in Worship Songs. "
                "A title stays on this list only when at least two distinct cached phrases "
                "were heard in at least two 30-second windows, or the match was only one step weaker than that."
            )
        ]
    )
    _style_sheet(song_sheet, [42, 36, 46, 18, 40, 16])
    for row in song_sheet.iter_rows(min_row=2, max_row=max(1, song_sheet.max_row - 2)):
        _link(row[2], row[2].value)

    topic_sheet = workbook.create_sheet("Topics & Verses")
    topic_sheet.append(TOPIC_HEADERS)
    by_theme: dict[str, list[dict]] = defaultdict(list)
    for record in sermons:
        by_theme[record["info"].get("theme") or "Unclassified"].append(record)
    theme_rows = []
    for theme in sorted(by_theme, key=lambda name: (-len(by_theme[name]), name)):
        group = by_theme[theme]
        text, count = _verse_text(group, index)
        dates = sorted({as_date(record["video"]["upload_date"]) for record in group}, reverse=True)
        mains = Counter(
            (record["info"].get("main_passage") or "").split(";")[0].split(":")[0].strip()
            for record in group
            if record["info"].get("main_passage")
        )
        main_bit = ", ".join(f"{name} ({times})" for name, times in mains.most_common(4))
        topic_sheet.append(
            [
                "THEME",
                theme,
                f"{len(group)} sermon(s). Most common main texts: {main_bit}",
                ", ".join(dates),
                text,
                count,
            ]
        )
        theme_rows.append(topic_sheet.max_row)
        for record in sorted(group, key=lambda item: as_date(item["video"]["upload_date"]), reverse=True):
            line, line_count = _verse_text([record], index)
            topic_sheet.append(
                [
                    "  sermon",
                    "   " + (record["info"].get("message") or theme),
                    record["info"].get("summary_en") or "",
                    as_date(record["video"]["upload_date"]),
                    line,
                    line_count,
                ]
            )
    _style_sheet(topic_sheet, [12, 48, 60, 18, 80, 14])
    for row_index in theme_rows:
        for cell in topic_sheet[row_index]:
            cell.font = Font(bold=True)
            cell.fill = THEME_FILL

    workbook.save(output)
    if export_csv:
        _write_csvs(output, workbook)
    return output


def _write_csvs(output: Path, workbook: Workbook) -> None:
    for worksheet in workbook.worksheets:
        slug = re.sub(r"[^A-Za-z0-9]+", "_", worksheet.title).strip("_").lower()
        path = output.with_name(f"{output.stem}_{slug}.csv")
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            for row in worksheet.iter_rows(values_only=True):
                if row is None or all(cell is None for cell in row):
                    continue
                writer.writerow(["" if cell is None else cell for cell in row])
